"""Gera um Excel com o diagnóstico de lacunas Fev-Ago/2026 para as 8 pastas de
`3.1 Registros Diarios MPFM` (B03, B08, B13, B05/PE-4, B10/PE-2, B15,
3.1.7_ALARMES_FCS_320 e 3.1.8_EVENTOS_FCS320).

Script somente leitura: não escreve no banco, não importa nenhum arquivo.
Não altera routes/system_routes.py -- reproduz (copiando, não chamando) a
mesma lógica de `_scan_month_source_coverage` e seus helpers, parametrizando
o que lá é capturado por closure (`monitor_prefs`, `db_conn`).

Uso:
    python scripts/gerar_relatorio_lacunas_fev_ago_2026.py
"""
from __future__ import annotations

import calendar
import glob
import json
import os
import re
import sqlite3
import sys
from datetime import datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app_config import DB_PATH, WORK_DIR

MONTHS = [f"2026-{m:02d}" for m in range(2, 9)]  # 2026-02 .. 2026-08

_B05_OLD_FOLDER = "3.1.4_18-FT-1506 PE 4 - Subsea B05"
_B05_OFFICIAL_FOLDER = "3.1.4_18-FT-1506 PE 4 e PE_EO105 - Subsea B05"

_PT_MONTH_NAMES = {
    1: "01. Janeiro", 2: "02. Fevereiro", 3: "03. Março", 4: "04. Abril",
    5: "05. Maio", 6: "06. Junho", 7: "07. Julho", 8: "08. Agosto",
    9: "09. Setembro", 10: "10. Outubro", 11: "11. Novembro", 12: "12. Dezembro",
}

SHEET_NAMES = {
    "B03": "3.1.1_B03",
    "B08": "3.1.2_B08",
    "B13": "3.1.3_B13",
    "B05": "3.1.4_B05_PE4",
    "B10": "3.1.5_B10_PE2",
    "B15": "3.1.6_B15",
}


def db_conn() -> sqlite3.Connection:
    db_target = str(DB_PATH)
    conn = sqlite3.connect(db_target, uri=db_target.startswith("file:"), timeout=30.0)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA busy_timeout = 30000")
    return conn


def _load_monitor_folders() -> list[dict]:
    """Reproduz a leitura de `_load_monitor_prefs` (routes/system_routes.py),
    somente lendo o arquivo -- não escreve nem recarrega estado do servidor."""
    prefs_path = WORK_DIR / "user_prefs.json"
    if not prefs_path.exists():
        return []
    try:
        payload = json.loads(prefs_path.read_text("utf-8"))
    except Exception:
        return []
    monitor = (payload or {}).get("auto_folder_monitor") or {}
    folders = monitor.get("folders") or []
    return [f for f in folders if isinstance(f, dict) and f.get("active", True)]


def _canonical_monitor_folder_path(folder_path: str) -> str:
    return str(folder_path or "").strip().replace(_B05_OLD_FOLDER, _B05_OFFICIAL_FOLDER)


def _discover_measurement_files(folder: str | Path) -> list[Path]:
    base = Path(folder).resolve()
    reserved_dirs = {"alarme", "alarmes", "alarm", "alarms"}
    all_files = sorted(glob.glob(os.path.join(str(base), "**", "*.pdf"), recursive=True)) + sorted(
        glob.glob(os.path.join(str(base), "**", "*.txt"), recursive=True)
    )
    candidates: list[Path] = []
    for path_str in all_files:
        path = Path(path_str)
        try:
            relative = path.resolve().relative_to(base)
        except ValueError:
            relative = path
        if any(part.lower() in reserved_dirs for part in relative.parts[:-1]):
            continue
        candidates.append(path)
    return candidates


def _parse_mpfm_pdf_filename(path: Path) -> dict | None:
    match = re.search(r"(?P<bank>B\d{2})_MPFM_(?P<kind>Daily|Hourly)-(?P<date>\d{8})-(?P<time>\d{6})", path.name, re.IGNORECASE)
    if not match:
        return None
    kind = match.group("kind").lower()
    try:
        end_dt = datetime.strptime(f'{match.group("date")}{match.group("time")}', "%Y%m%d%H%M%S")
    except ValueError:
        return None
    if kind == "daily":
        production_day = (end_dt.date() - timedelta(days=1)).isoformat()
        hour = None
    else:
        start_dt = end_dt - timedelta(hours=1)
        production_day = start_dt.date().isoformat()
        hour = start_dt.hour
    return {
        "path": path,
        "name": path.name,
        "bank": match.group("bank").upper(),
        "report_type": kind,
        "production_day": production_day,
        "hour": hour,
    }


def _candidate_existing_mpfm_roots(folder_path: str) -> list[Path]:
    raw = Path(_canonical_monitor_folder_path(folder_path))
    variants: list[Path] = []
    if str(raw):
        variants.append(raw)
        parts = list(raw.parts)
        for idx, part in enumerate(parts):
            if part == "3. Registros de Operação SGM Multifasico" and idx + 1 < len(parts) and parts[idx + 1] == "3.1 Registros Diarios MPFM":
                try:
                    variants.append(Path(parts[0], *parts[1:idx], *parts[idx + 1:]))
                except Exception:
                    pass
                break

    roots: list[Path] = []
    seen: set[str] = set()
    month_match = re.match(r"^(?P<num>\d{2})\.", raw.name or "")
    preferred_months: set[int] = set()
    if month_match:
        try:
            month_num = int(month_match.group("num"))
            preferred_months.update({month_num - 1, month_num, month_num + 1})
        except Exception:
            preferred_months = set()
    for variant in variants:
        if not variant.exists() or not variant.is_dir():
            continue
        scan_roots = [variant]
        if re.match(r"^\d{2}\.", variant.name or "") and variant.parent.is_dir() and re.match(r"^\d{4}$", variant.parent.name or ""):
            siblings = [item for item in variant.parent.iterdir() if item.is_dir() and re.match(r"^\d{2}\.", item.name or "")]
            filtered = []
            for item in siblings:
                item_match = re.match(r"^(\d{2})\.", item.name or "")
                if not item_match:
                    continue
                try:
                    if int(item_match.group(1)) in preferred_months:
                        filtered.append(item)
                except Exception:
                    continue
            scan_roots = filtered or [variant]
        for root in scan_roots:
            key = str(root.resolve()).casefold()
            if key in seen:
                continue
            seen.add(key)
            roots.append(root.resolve())
    return sorted(roots, key=lambda item: str(item).casefold())


def _discover_mpfm_pdf_files_fast(scan_root: Path) -> list[Path]:
    candidates: list[Path] = []
    if scan_root.name.lower() in {"daily", "hourly"}:
        candidates.extend(scan_root.glob("*.pdf"))
    else:
        seen_dirs: set[str] = set()
        for child_name in ("Daily", "DAILY", "daily", "Hourly", "HOURLY", "hourly"):
            child = scan_root / child_name
            dir_key = str(child).casefold()
            if dir_key in seen_dirs:
                continue
            seen_dirs.add(dir_key)
            if child.is_dir():
                candidates.extend(child.glob("*.pdf"))
        candidates.extend(scan_root.glob("*.pdf"))
    seen: set[str] = set()
    result: list[Path] = []
    for path in candidates:
        if not path.is_file():
            continue
        resolved = path.resolve()
        key = str(resolved).casefold()
        if key in seen:
            continue
        seen.add(key)
        result.append(resolved)
    return sorted(result, key=lambda item: item.name.casefold())


def _meter_root_from_monitor_path(folder_path: str) -> Path | None:
    path = Path(_canonical_monitor_folder_path(folder_path))
    parts = list(path.parts)
    if len(parts) >= 2 and re.match(r"^\d{2}\.", parts[-1] or "") and re.match(r"^\d{4}$", parts[-2] or ""):
        return Path(*parts[:-2])
    return None


def scan_month_source_coverage(month: str, monitor_folders: list[dict]) -> dict:
    """Cópia parametrizada de `_scan_month_source_coverage`
    (routes/system_routes.py:476-629): mesma lógica, recebendo
    `monitor_folders` como argumento em vez de closure sobre `monitor_prefs`."""
    match = re.match(r"^(?P<year>\d{4})-(?P<mon>\d{2})$", str(month or "").strip())
    if not match:
        raise ValueError("Mês inválido. Use o formato YYYY-MM.")
    year = int(match.group("year"))
    mon = int(match.group("mon"))
    month_name = _PT_MONTH_NAMES.get(mon)
    if not month_name:
        raise ValueError("Mês inválido.")

    days_in_month = calendar.monthrange(year, mon)[1]
    month_start = f"{year:04d}-{mon:02d}-01"
    month_end = f"{year:04d}-{mon:02d}-{days_in_month:02d}"

    conn = db_conn()
    cur = conn.cursor()
    rows = cur.execute(
        "SELECT bank, day_ref, row_kind, hour_ref FROM measurements_curated "
        "WHERE row_kind IN ('daily','hourly') AND bank<>'' AND bank<>'SEP' "
        "AND day_ref BETWEEN ? AND ?",
        (month_start, month_end),
    ).fetchall()
    conn.close()

    db_daily: set[tuple[str, str]] = set()
    db_hours: dict[tuple[str, str], set[int]] = {}
    for bank, day_ref, row_kind, hour_ref in rows:
        if row_kind == "daily":
            db_daily.add((bank, day_ref))
        elif row_kind == "hourly" and hour_ref is not None:
            db_hours.setdefault((bank, day_ref), set()).add(int(hour_ref))

    today = datetime.now().date()
    last_day = min(days_in_month, today.day) if (today.year == year and today.month == mon) else days_in_month
    all_month_days = [f"{year:04d}-{mon:02d}-{day:02d}" for day in range(1, last_day + 1)]

    banks_out: list[dict] = []

    for folder in monitor_folders:
        folder_path = str(folder.get("path") or "").strip()
        label = folder.get("label") or Path(folder_path).name
        bank_match = re.search(r"\bB\d{2}\b", folder_path, re.IGNORECASE)
        expected_bank = bank_match.group(0).upper() if bank_match else ""
        bank_entry = {
            "bank": expected_bank,
            "label": label,
            "folder_path": folder_path,
            "scan_roots": [],
            "found_on_disk": False,
            "pending_count": 0,
            "source_gap_count": 0,
            "days": [],
            "error": "",
        }
        if not expected_bank:
            bank_entry["error"] = "Não foi possível identificar o código do banco (B##) na pasta configurada."
            banks_out.append(bank_entry)
            continue
        meter_root = _meter_root_from_monitor_path(folder_path)
        if meter_root is None:
            bank_entry["error"] = f"Estrutura de pasta inesperada (esperado .../{{ano}}/{{mês}}): {folder_path}"
            banks_out.append(bank_entry)
            continue

        synthetic_path = str(meter_root / f"{year:04d}" / month_name)
        scan_roots = _candidate_existing_mpfm_roots(synthetic_path)
        bank_entry["scan_roots"] = [str(item) for item in scan_roots]
        bank_entry["found_on_disk"] = bool(scan_roots)
        if not scan_roots:
            bank_entry["error"] = f"Pasta do mês não encontrada em disco: {synthetic_path}"

        disk_daily: dict[str, Path] = {}
        disk_hours: dict[str, dict[int, Path]] = {}
        files: list[Path] = []
        for scan_root in scan_roots:
            files.extend(_discover_mpfm_pdf_files_fast(scan_root))
        if not files and scan_roots:
            for scan_root in scan_roots:
                files.extend(_discover_measurement_files(scan_root))
        for path in files:
            meta = _parse_mpfm_pdf_filename(path)
            if not meta or meta["bank"] != expected_bank:
                continue
            day = meta["production_day"]
            if not day.startswith(f"{year:04d}-{mon:02d}"):
                continue
            if meta["report_type"] == "daily":
                disk_daily[day] = meta["path"]
            elif meta.get("hour") is not None:
                disk_hours.setdefault(day, {})[meta["hour"]] = meta["path"]

        day_reports = []
        bank_pending = 0
        bank_gap = 0
        for day in all_month_days:
            key = (expected_bank, day)
            missing_hours: list[int] = []
            pending_hours: list[dict] = []
            for hour in range(24):
                if hour in db_hours.get(key, set()):
                    continue
                disk_path = disk_hours.get(day, {}).get(hour)
                if disk_path is not None:
                    pending_hours.append({"hour": hour, "path": str(disk_path)})
                else:
                    missing_hours.append(hour)
            daily_pending = False
            daily_missing = False
            if key not in db_daily:
                disk_daily_path = disk_daily.get(day)
                if disk_daily_path is not None:
                    daily_pending = True
                else:
                    daily_missing = True
            if missing_hours or pending_hours or daily_pending or daily_missing:
                day_reports.append(
                    {
                        "day": day,
                        "missing_hours": missing_hours,
                        "pending_hours": pending_hours,
                        "daily_missing": daily_missing,
                        "daily_pending": daily_pending,
                    }
                )
            bank_pending += len(pending_hours) + (1 if daily_pending else 0)
            bank_gap += len(missing_hours) + (1 if daily_missing else 0)

        bank_entry["pending_count"] = bank_pending
        bank_entry["source_gap_count"] = bank_gap
        bank_entry["days"] = day_reports
        banks_out.append(bank_entry)

    return {"month": f"{year:04d}-{mon:02d}", "banks": banks_out}


# --- 3.1.7 ALARMES_FCS_320 (comparação por nome de arquivo) ---------------

ONEDRIVE_MPFM_ROOT = Path(
    r"C:\Users\MAUAM\OneDrive - Equinor\DPB FPSO Bacalhau - Metering - 02 MULTIPHASE MANAGEMENT SYSTEM"
    r"\3.1 Registros Diarios MPFM"
)
ALARMES_FCS320_ROOT = ONEDRIVE_MPFM_ROOT / "3.1.7_ALARMES_FCS_320" / "2026"
EVENTOS_FCS320_ROOT = ONEDRIVE_MPFM_ROOT / "3.1.8_EVENTOS_FCS320" / "2026"


def _norm(name: str) -> str:
    return re.sub(r"[\s_]+", "_", name.strip().lower())


def _imported_alarm_basenames() -> set[str]:
    conn = db_conn()
    try:
        refs = [r[0] for r in conn.execute(
            "SELECT DISTINCT source_ref FROM alarm_records WHERE source_kind='pdf'"
        ).fetchall()]
    finally:
        conn.close()
    seen = set()
    for ref in refs:
        for part in ref.split(";"):
            part = part.strip()
            if part.lower().startswith("pdf:"):
                part = part[4:]
            name = Path(part).name if ("\\" in part or "/" in part) else part
            if name.lower().endswith(".pdf"):
                seen.add(_norm(name))
    return seen


def scan_alarmes_fcs320(months: list[str]) -> list[dict]:
    imported = _imported_alarm_basenames()
    month_nums = {int(m.split("-")[1]) for m in months}
    rows: list[dict] = []
    if not ALARMES_FCS320_ROOT.exists():
        return rows
    for month_dir in sorted(ALARMES_FCS320_ROOT.iterdir()):
        month_match = re.match(r"^(\d{2})\.", month_dir.name or "")
        if not month_dir.is_dir() or not month_match or int(month_match.group(1)) not in month_nums:
            continue
        for pdf in sorted(month_dir.glob("*.pdf")):
            day_match = re.match(r"^(\d{2})-(\d{2})_", pdf.name)
            production_day = (
                f"2026-{day_match.group(2)}-{day_match.group(1)}" if day_match else ""
            )
            rows.append(
                {
                    "arquivo": pdf.name,
                    "pasta_mes": month_dir.name,
                    "dia_producao": production_day,
                    "importado": _norm(pdf.name) in imported,
                }
            )
    return rows


def scan_eventos_fcs320(months: list[str]) -> list[dict]:
    month_nums = {int(m.split("-")[1]) for m in months}
    rows: list[dict] = []
    if not EVENTOS_FCS320_ROOT.exists():
        return rows
    for month_dir in sorted(EVENTOS_FCS320_ROOT.iterdir()):
        month_match = re.match(r"^(\d{2})\.", month_dir.name or "")
        if not month_dir.is_dir() or not month_match or int(month_match.group(1)) not in month_nums:
            continue
        files = sorted(month_dir.glob("*.pdf")) + sorted(month_dir.glob("*.txt"))
        for f in files:
            try:
                mtime = datetime.fromtimestamp(f.stat().st_mtime).date().isoformat()
            except OSError:
                mtime = ""
            rows.append(
                {
                    "arquivo": f.name,
                    "pasta_mes": month_dir.name,
                    "data_modificacao": mtime,
                    "status": "Sem pipeline de importação (lacuna de funcionalidade, não de dado)",
                }
            )
    return rows


# --- Montagem do Excel ------------------------------------------------------

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFFFF", bold=True)
GAP_FILL = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
PENDING_FILL = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")


def _write_header(ws, headers: list[str]) -> None:
    for col_idx, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"


def _autosize(ws, headers: list[str]) -> None:
    for col_idx, title in enumerate(headers, start=1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = max(14, min(60, len(title) + 4))


def build_bank_sheet(wb: Workbook, sheet_name: str, month_results: list[dict]) -> None:
    ws = wb.create_sheet(sheet_name)
    headers = ["Mês", "Banco", "Pasta Monitorada", "Dia", "Horas Faltantes (lacuna de origem)",
               "Horas Pendentes de Importação", "Diário Faltante", "Diário Pendente", "Erro/Observação"]
    _write_header(ws, headers)
    row_idx = 2
    for result in month_results:
        month = result["month"]
        for bank_entry in result["banks"]:
            if bank_entry.get("error") and not bank_entry.get("days"):
                ws.cell(row=row_idx, column=1, value=month)
                ws.cell(row=row_idx, column=2, value=bank_entry["bank"])
                ws.cell(row=row_idx, column=3, value=bank_entry["folder_path"])
                ws.cell(row=row_idx, column=9, value=bank_entry["error"])
                row_idx += 1
                continue
            for day_report in bank_entry["days"]:
                ws.cell(row=row_idx, column=1, value=month)
                ws.cell(row=row_idx, column=2, value=bank_entry["bank"])
                ws.cell(row=row_idx, column=3, value=bank_entry["folder_path"])
                ws.cell(row=row_idx, column=4, value=day_report["day"])
                missing_cell = ws.cell(row=row_idx, column=5, value=", ".join(str(h) for h in day_report["missing_hours"]) or "")
                pending_cell = ws.cell(row=row_idx, column=6, value=", ".join(str(p["hour"]) for p in day_report["pending_hours"]) or "")
                ws.cell(row=row_idx, column=7, value="SIM" if day_report["daily_missing"] else "")
                ws.cell(row=row_idx, column=8, value="SIM" if day_report["daily_pending"] else "")
                if day_report["missing_hours"]:
                    missing_cell.fill = GAP_FILL
                if day_report["daily_missing"]:
                    ws.cell(row=row_idx, column=7).fill = GAP_FILL
                if day_report["pending_hours"]:
                    pending_cell.fill = PENDING_FILL
                if day_report["daily_pending"]:
                    ws.cell(row=row_idx, column=8).fill = PENDING_FILL
                row_idx += 1
    _autosize(ws, headers)


def build_alarmes_sheet(wb: Workbook, rows: list[dict]) -> None:
    ws = wb.create_sheet("3.1.7_ALARMES_FCS320")
    headers = ["Arquivo", "Pasta (mês)", "Dia de produção", "Já importado?"]
    _write_header(ws, headers)
    for row_idx, row in enumerate(rows, start=2):
        ws.cell(row=row_idx, column=1, value=row["arquivo"])
        ws.cell(row=row_idx, column=2, value=row["pasta_mes"])
        ws.cell(row=row_idx, column=3, value=row["dia_producao"])
        cell = ws.cell(row=row_idx, column=4, value="SIM" if row["importado"] else "NÃO")
        if not row["importado"]:
            cell.fill = GAP_FILL
    _autosize(ws, headers)


def build_eventos_sheet(wb: Workbook, rows: list[dict]) -> None:
    ws = wb.create_sheet("3.1.8_EVENTOS_FCS320")
    headers = ["Arquivo", "Pasta (mês)", "Data de modificação", "Status"]
    _write_header(ws, headers)
    for row_idx, row in enumerate(rows, start=2):
        ws.cell(row=row_idx, column=1, value=row["arquivo"])
        ws.cell(row=row_idx, column=2, value=row["pasta_mes"])
        ws.cell(row=row_idx, column=3, value=row["data_modificacao"])
        cell = ws.cell(row=row_idx, column=4, value=row["status"])
        cell.fill = PENDING_FILL
    _autosize(ws, headers)


def build_summary_sheet(wb: Workbook, month_results: list[dict], alarmes_rows: list[dict], eventos_rows: list[dict]) -> None:
    ws = wb.create_sheet("Resumo", 0)
    headers = ["Pasta", "Aba", "Total Lacuna de Origem", "Total Pendente de Importação", "Observação"]
    _write_header(ws, headers)
    row_idx = 2
    for bank, sheet_name in SHEET_NAMES.items():
        total_gap = 0
        total_pending = 0
        for result in month_results:
            for bank_entry in result["banks"]:
                if bank_entry["bank"] == bank:
                    total_gap += bank_entry["source_gap_count"]
                    total_pending += bank_entry["pending_count"]
        ws.cell(row=row_idx, column=1, value=bank)
        ws.cell(row=row_idx, column=2, value=sheet_name)
        ws.cell(row=row_idx, column=3, value=total_gap)
        ws.cell(row=row_idx, column=4, value=total_pending)
        row_idx += 1
    total_not_imported = sum(1 for r in alarmes_rows if not r["importado"])
    ws.cell(row=row_idx, column=1, value="ALARMES_FCS_320")
    ws.cell(row=row_idx, column=2, value="3.1.7_ALARMES_FCS320")
    ws.cell(row=row_idx, column=3, value=0)
    ws.cell(row=row_idx, column=4, value=total_not_imported)
    ws.cell(row=row_idx, column=5, value=f"{len(alarmes_rows)} PDFs em disco (fev-ago/2026)")
    row_idx += 1
    ws.cell(row=row_idx, column=1, value="EVENTOS_FCS320")
    ws.cell(row=row_idx, column=2, value="3.1.8_EVENTOS_FCS320")
    ws.cell(row=row_idx, column=3, value=0)
    ws.cell(row=row_idx, column=4, value=len(eventos_rows))
    ws.cell(row=row_idx, column=5, value="Sem pipeline de importação (lacuna de funcionalidade)")
    _autosize(ws, headers)


def main() -> None:
    print(f"Meses analisados: {MONTHS[0]} a {MONTHS[-1]}")

    monitor_folders = _load_monitor_folders()
    by_bank: dict[str, list[dict]] = {bank: [] for bank in SHEET_NAMES}
    for folder in monitor_folders:
        bank_match = re.search(r"\bB\d{2}\b", str(folder.get("path") or ""), re.IGNORECASE)
        bank = bank_match.group(0).upper() if bank_match else ""
        if bank in by_bank:
            by_bank[bank].append(folder)

    month_results_by_bank: dict[str, list[dict]] = {bank: [] for bank in SHEET_NAMES}
    for month in MONTHS:
        print(f"Varrendo mês {month}...", flush=True)
        for bank, folders in by_bank.items():
            if not folders:
                continue
            result = scan_month_source_coverage(month, folders)
            month_results_by_bank[bank].append(result)

    print("Varrendo 3.1.7_ALARMES_FCS_320...", flush=True)
    alarmes_rows = scan_alarmes_fcs320(MONTHS)
    print("Varrendo 3.1.8_EVENTOS_FCS320...", flush=True)
    eventos_rows = scan_eventos_fcs320(MONTHS)

    wb = Workbook()
    wb.remove(wb.active)

    all_results_flat = [r for results in month_results_by_bank.values() for r in results]
    build_summary_sheet(wb, all_results_flat, alarmes_rows, eventos_rows)

    for bank, sheet_name in SHEET_NAMES.items():
        build_bank_sheet(wb, sheet_name, month_results_by_bank[bank])

    build_alarmes_sheet(wb, alarmes_rows)
    build_eventos_sheet(wb, eventos_rows)

    out_dir = ROOT / "AUTOMACAO_HTML" / "HTML_GERADOS"
    out_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = out_dir / f"RELATORIO_LACUNAS_FEV_AGO_2026_{timestamp}.xlsx"
    wb.save(out_path)
    print(f"\nRelatório salvo em: {out_path}")


if __name__ == "__main__":
    main()
