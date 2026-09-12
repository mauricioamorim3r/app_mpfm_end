import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from services.excel_template_service import MONTHLY_WORKBOOK_TEMPLATE
from services.importing.monthly_workbook_service import (
    BASE_UNICA_COLUMNS,
    _replace_sheet_in_workbook,
    _write_fast_monthly_workbook,
)


def test_replace_base_unica_keeps_template_format_without_stale_rows():
    template = load_workbook(MONTHLY_WORKBOOK_TEMPLATE)
    workbook = load_workbook(MONTHLY_WORKBOOK_TEMPLATE)
    frame = pd.DataFrame(
        [
            {
                "ProductionDate": "2026-08-01",
                "Hour": "00:00",
                "Granularity": "Hourly",
                "Origin": "MPFM",
                "Bank": "B03",
            }
        ],
        columns=BASE_UNICA_COLUMNS,
    ).fillna("")

    _replace_sheet_in_workbook(workbook, template, "BASE_UNICA_MES", frame)

    sheet = workbook["BASE_UNICA_MES"]
    assert sheet.max_row == 3
    assert sheet.cell(2, 2).value == "2026-08-01"
    assert all(sheet.cell(3, column).value is None for column in range(2, 2 + len(BASE_UNICA_COLUMNS)))
    assert sheet.cell(2, 2).style_id == template["BASE_UNICA_MES"].cell(2, 2).style_id
    last_column = get_column_letter(1 + len(BASE_UNICA_COLUMNS))
    assert sheet.tables["tblBaseUnicaMes"].ref == f"B1:{last_column}2"
    workbook.close()
    template.close()


def test_fast_workbook_preserves_sheet_contract(tmp_path):
    base = pd.DataFrame(
        [{"ProductionDate": "2026-08-01", "Bank": "B03"}],
        columns=BASE_UNICA_COLUMNS,
    ).fillna("")
    alerts = pd.DataFrame(
        [["23_08", "warn", "missing_hours", "B03_23_08", "Horas faltando: 23", "2026-09-01"]],
        columns=["day_ref", "severity", "issue_type", "ref_key", "details", "created_at"],
    )
    status = pd.DataFrame([["ok"]], columns=["Status"])
    output = tmp_path / "month.xlsx"

    _write_fast_monthly_workbook(
        output,
        [
            ("BASE_UNICA_MES", base, 1),
            ("DAILYS", base, 2),
            ("HOURLYS", base, 2),
            ("RECON", base, 2),
            ("ALERTAS_MES", alerts, 2),
            ("STATUS_MES", status, 0),
        ],
    )

    workbook = load_workbook(output, read_only=False, data_only=True)
    assert workbook.sheetnames == [
        "BASE_UNICA_MES",
        "DAILYS",
        "HOURLYS",
        "RECON",
        "ALERTAS_MES",
        "STATUS_MES",
    ]
    assert workbook["BASE_UNICA_MES"].tables["tblBaseUnicaMes"]
    assert workbook["DAILYS"]["A1"].value == "NAVEGAÇÃO"
    assert workbook["STATUS_MES"]["A2"].value == "ok"
    workbook.close()
